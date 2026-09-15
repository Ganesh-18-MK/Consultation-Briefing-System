"""One-time migration: insert a new 'Owner' column between Email and the
first date column in the leads sheet, so mam/staff can manually assign a
client to a staff member. Existing date columns and all client rows are
preserved untouched — only a single blank column is inserted at index 3
with the header 'Owner'. Safe to run more than once: if column C is
already 'Owner', it exits without touching the file.

Usage (run against a LOCAL copy downloaded from the bucket — see the
surrounding terminal commands for the download/upload steps):
    python3 scripts/migrate_add_owner_column.py /path/to/leads.xlsx
"""
import sys
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    if len(sys.argv) != 2:
        print("Usage: python3 scripts/migrate_add_owner_column.py <path-to-leads.xlsx>")
        sys.exit(1)

    path = Path(sys.argv[1])
    wb = load_workbook(path)
    ws = wb.active

    existing_c1 = ws.cell(row=1, column=3).value
    if existing_c1 == "Owner":
        print("Column C is already 'Owner' — nothing to do.")
        return

    ws.insert_cols(3)
    ws.cell(row=1, column=3, value="Owner")
    wb.save(path)
    print(f"Inserted a blank 'Owner' column at C in {path}")


if __name__ == "__main__":
    main()
