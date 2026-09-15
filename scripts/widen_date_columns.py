"""One-time migration: widen every existing date column (D onward) to
match the new default width used for columns created from now on (see
app/leads_sheet.py's _find_or_create_date_column). Only touches column
width, not any cell content — safe to re-run.

Usage (run against a LOCAL copy downloaded from the bucket):
    python3 scripts/widen_date_columns.py /path/to/leads.xlsx
"""
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

NEW_WIDTH = 100
FIRST_DATE_COL = 4  # column D — keep in sync with app/leads_sheet.py


def main() -> None:
    if len(sys.argv) != 2:
        print("Usage: python3 scripts/widen_date_columns.py <path-to-leads.xlsx>")
        sys.exit(1)

    path = Path(sys.argv[1])
    wb = load_workbook(path)
    ws = wb.active

    changed = 0
    for col_idx in range(FIRST_DATE_COL, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        current = ws.column_dimensions[letter].width
        if current != NEW_WIDTH:
            ws.column_dimensions[letter].width = NEW_WIDTH
            changed += 1

    wb.save(path)
    print(f"Widened {changed} date column(s) to width {NEW_WIDTH} in {path}")


if __name__ == "__main__":
    main()
