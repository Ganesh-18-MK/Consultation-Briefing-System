from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import Workbook, load_workbook

from app import db, leads_sheet
from app.config import settings


def _make_booking(start_iso=None):
    client_id, _ = db.upsert_client("jane@example.com", "Jane Client")
    start = start_iso or (datetime.now(timezone.utc) + timedelta(minutes=10)).isoformat()
    booking_id = db.create_booking(
        "evt-1", client_id, "attorney@example.com", start, "Wants an H-1B extension", None, False
    )
    return booking_id


def test_format_date_time_splits_iso_string():
    # 15:30 UTC -> 10:30 AM Central (September is Daylight Time, UTC-5),
    # same calendar day.
    date_str, time_str = leads_sheet.format_date_time("2026-09-10T15:30:00+00:00")
    assert date_str == "2026-09-10"
    assert time_str == "10:30 AM CDT"


def test_format_date_time_rolls_date_backward_across_midnight_ct():
    # 03:00 UTC -> 10:00 PM Central the *previous* calendar day — Central
    # Time is behind UTC, so an early-UTC-morning booking rolls the date
    # backward (unlike a timezone ahead of UTC, which would roll forward).
    # Date and Time must agree with each other (both derived from the
    # same Central-Time datetime).
    date_str, time_str = leads_sheet.format_date_time("2026-09-12T03:00:00+00:00")
    assert date_str == "2026-09-11"
    assert time_str == "10:00 PM CDT"


def test_format_date_time_handles_missing_value():
    assert leads_sheet.format_date_time(None) == ("", "")


def test_format_date_time_handles_unparseable_value():
    date_str, time_str = leads_sheet.format_date_time("not-a-date")
    assert date_str == "not-a-date"
    assert time_str == ""


def test_format_date_splits_iso_string():
    assert leads_sheet.format_date("2026-09-10T15:30:00+00:00") == "2026-09-10"


def test_format_date_handles_missing_value():
    assert leads_sheet.format_date(None) == ""


def test_format_date_handles_unparseable_value():
    assert leads_sheet.format_date("not-a-date") == "not-a-date"


def test_creates_fresh_file_with_fixed_headers_when_missing(tmp_path, monkeypatch):
    dest = tmp_path / "subdir" / "leads.xlsx"
    monkeypatch.setattr(settings, "leads_sheet_path", str(dest))
    monkeypatch.setattr(settings, "leads_sheet_enabled", True)

    leads_sheet.record_client_response("Jane Client", "jane@example.com", "2026-09-10", "H-1B extension")

    assert dest.exists()
    wb = load_workbook(dest)
    ws = wb.active
    header_row = [cell.value for cell in ws[1]]
    assert header_row == ["Client Name", "Email", "Owner", "2026-09-10"]
    data_row = [cell.value for cell in ws[2]]
    assert data_row == ["Jane Client", "jane@example.com", None, "H-1B extension"]
    assert list(dest.parent.glob("*.tmp")) == []


def test_new_client_new_date_gets_new_row_and_new_column(tmp_path, monkeypatch):
    dest = tmp_path / "leads.xlsx"
    monkeypatch.setattr(settings, "leads_sheet_path", str(dest))
    monkeypatch.setattr(settings, "leads_sheet_enabled", True)

    leads_sheet.record_client_response("Jane Client", "jane@example.com", "2026-09-10", "H-1B extension")
    leads_sheet.record_client_response("John Other", "john@example.com", "2026-09-11", "Green card")

    wb = load_workbook(dest)
    ws = wb.active
    assert [cell.value for cell in ws[1]] == ["Client Name", "Email", "Owner", "2026-09-10", "2026-09-11"]
    assert [cell.value for cell in ws[2]] == ["Jane Client", "jane@example.com", None, "H-1B extension", None]
    assert [cell.value for cell in ws[3]] == ["John Other", "john@example.com", None, None, "Green card"]


def test_repeat_client_new_date_reuses_row_adds_column(tmp_path, monkeypatch):
    dest = tmp_path / "leads.xlsx"
    monkeypatch.setattr(settings, "leads_sheet_path", str(dest))
    monkeypatch.setattr(settings, "leads_sheet_enabled", True)

    leads_sheet.record_client_response("Jane Client", "jane@example.com", "2026-09-10", "H-1B extension")
    # Same client (matched by email), booking again on a new date.
    leads_sheet.record_client_response("Jane Client", "jane@example.com", "2026-10-01", "PERM timeline question")

    wb = load_workbook(dest)
    ws = wb.active
    assert ws.max_row == 2  # still just one row for Jane, not a second one
    assert [cell.value for cell in ws[1]] == ["Client Name", "Email", "Owner", "2026-09-10", "2026-10-01"]
    assert [cell.value for cell in ws[2]] == [
        "Jane Client",
        "jane@example.com",
        None,
        "H-1B extension",
        "PERM timeline question",
    ]


def test_two_different_clients_same_date_share_the_column(tmp_path, monkeypatch):
    dest = tmp_path / "leads.xlsx"
    monkeypatch.setattr(settings, "leads_sheet_path", str(dest))
    monkeypatch.setattr(settings, "leads_sheet_enabled", True)

    leads_sheet.record_client_response("Jane Client", "jane@example.com", "2026-09-10", "H-1B extension")
    leads_sheet.record_client_response("John Other", "john@example.com", "2026-09-10", "Green card")

    wb = load_workbook(dest)
    ws = wb.active
    assert [cell.value for cell in ws[1]] == ["Client Name", "Email", "Owner", "2026-09-10"]
    assert ws.max_row == 3


def test_client_matched_by_email_case_insensitively(tmp_path, monkeypatch):
    dest = tmp_path / "leads.xlsx"
    monkeypatch.setattr(settings, "leads_sheet_path", str(dest))
    monkeypatch.setattr(settings, "leads_sheet_enabled", True)

    leads_sheet.record_client_response("Jane Client", "Jane@Example.com", "2026-09-10", "H-1B extension")
    leads_sheet.record_client_response("Jane Client", "jane@example.com", "2026-10-01", "Follow-up")

    wb = load_workbook(dest)
    ws = wb.active
    assert ws.max_row == 2


def test_cell_has_wrap_text_enabled_for_multiline_content(tmp_path, monkeypatch):
    dest = tmp_path / "leads.xlsx"
    monkeypatch.setattr(settings, "leads_sheet_path", str(dest))
    monkeypatch.setattr(settings, "leads_sheet_enabled", True)

    leads_sheet.record_client_response(
        "Jane Client", "jane@example.com", "2026-09-10", "1. First point.\n2. Second point."
    )

    wb = load_workbook(dest)
    ws = wb.active
    cell = ws.cell(row=2, column=4)
    assert cell.value == "1. First point.\n2. Second point."
    assert cell.alignment.wrap_text is True


def test_new_date_column_gets_wider_default_width(tmp_path, monkeypatch):
    dest = tmp_path / "leads.xlsx"
    monkeypatch.setattr(settings, "leads_sheet_path", str(dest))
    monkeypatch.setattr(settings, "leads_sheet_enabled", True)

    leads_sheet.record_client_response("Jane Client", "jane@example.com", "2026-09-10", "Some notes")

    wb = load_workbook(dest)
    ws = wb.active
    assert ws.column_dimensions["D"].width == 100


def test_skips_when_disabled(tmp_path, monkeypatch):
    dest = tmp_path / "leads.xlsx"
    monkeypatch.setattr(settings, "leads_sheet_path", str(dest))
    monkeypatch.setattr(settings, "leads_sheet_enabled", False)

    leads_sheet.record_client_response("Jane Client", "jane@example.com", "2026-09-10", "H-1B extension")

    assert not dest.exists()


def test_backfill_records_and_marks_synced(temp_db, tmp_path, monkeypatch):
    dest = tmp_path / "leads.xlsx"
    monkeypatch.setattr(settings, "leads_sheet_path", str(dest))
    monkeypatch.setattr(settings, "leads_sheet_enabled", True)

    _make_booking()

    leads_sheet.backfill_missing_rows()

    wb = load_workbook(dest)
    assert wb.active.max_row == 2

    pending = db.get_bookings_needing_leads_sheet_sync()
    assert pending == []
    # Re-running should not duplicate the row.
    leads_sheet.backfill_missing_rows()
    wb2 = load_workbook(dest)
    assert wb2.active.max_row == 2


def test_backfill_noop_when_nothing_pending(temp_db, tmp_path, monkeypatch):
    dest = tmp_path / "leads.xlsx"
    monkeypatch.setattr(settings, "leads_sheet_path", str(dest))
    monkeypatch.setattr(settings, "leads_sheet_enabled", True)

    leads_sheet.backfill_missing_rows()

    assert not dest.exists()
