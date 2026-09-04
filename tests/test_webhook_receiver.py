import hashlib
import hmac
import json
import time
from pathlib import Path

import pytest

from app import calendly_client, db
from app.config import settings
from app.webhook_receiver import app as flask_app

FIXTURE_PATH = Path(__file__).parent / "fixtures" / "calendly_invitee_created.json"


@pytest.fixture()
def client():
    flask_app.config["TESTING"] = True
    return flask_app.test_client()


@pytest.fixture()
def fake_event_details(monkeypatch):
    """Stub the Calendly API enrichment call so tests don't need network
    access or a real CALENDLY_API_TOKEN."""
    def _fake(event_uri):
        return {
            "start_time": "2026-09-02T18:30:00Z",
            "attorney_email": "attorney@example.com",
            "join_url": "https://teams.microsoft.com/l/meetup-join/fake",
        }

    monkeypatch.setattr(calendly_client, "get_event_details", _fake)


def _sign(body_bytes: bytes, secret: str) -> str:
    ts = str(int(time.time()))
    sig = hmac.new(secret.encode(), f"{ts}.".encode() + body_bytes, hashlib.sha256).hexdigest()
    return f"t={ts},v1={sig}"


def test_invitee_created_records_new_client(temp_db, no_signature_check, fake_event_details, client):
    body = json.loads(FIXTURE_PATH.read_text())
    resp = client.post("/webhooks/calendly", data=json.dumps(body), content_type="application/json")

    assert resp.status_code == 201
    booking = db.get_booking_by_uuid("INVITEE_UUID")
    assert booking is not None
    assert booking["attorney_email"] == "attorney@example.com"
    assert booking["is_repeat_client"] == 0
    assert "H-1B extension" in booking["discussion_notes"]


def test_second_booking_flagged_as_repeat_client(temp_db, no_signature_check, fake_event_details, client):
    body = json.loads(FIXTURE_PATH.read_text())
    client.post("/webhooks/calendly", data=json.dumps(body), content_type="application/json")

    # Same client, different event/invitee uuid.
    body2 = json.loads(FIXTURE_PATH.read_text())
    body2["payload"]["uri"] = "https://api.calendly.com/scheduled_events/EVENT_UUID_2/invitees/INVITEE_UUID_2"
    body2["payload"]["event"] = "https://api.calendly.com/scheduled_events/EVENT_UUID_2"
    client.post("/webhooks/calendly", data=json.dumps(body2), content_type="application/json")

    second_booking = db.get_booking_by_uuid("INVITEE_UUID_2")
    assert second_booking["is_repeat_client"] == 1


def test_duplicate_delivery_is_idempotent(temp_db, no_signature_check, fake_event_details, client):
    body = json.loads(FIXTURE_PATH.read_text())
    r1 = client.post("/webhooks/calendly", data=json.dumps(body), content_type="application/json")
    r2 = client.post("/webhooks/calendly", data=json.dumps(body), content_type="application/json")

    assert r1.status_code == 201
    assert r2.status_code == 200
    assert r2.get_json()["status"] == "duplicate"


def test_invitee_canceled_marks_booking_canceled(temp_db, no_signature_check, fake_event_details, client):
    body = json.loads(FIXTURE_PATH.read_text())
    client.post("/webhooks/calendly", data=json.dumps(body), content_type="application/json")

    cancel_body = {
        "event": "invitee.canceled",
        "payload": {"uri": body["payload"]["uri"]},
    }
    resp = client.post("/webhooks/calendly", data=json.dumps(cancel_body), content_type="application/json")

    assert resp.status_code == 200
    booking = db.get_booking_by_uuid("INVITEE_UUID")
    assert booking["status"] == "canceled"


def test_invalid_signature_rejected(temp_db, fake_event_details, client, monkeypatch):
    monkeypatch.setattr(settings, "calendly_signing_secret", "test-secret")
    body = json.loads(FIXTURE_PATH.read_text())
    body_bytes = json.dumps(body).encode()

    resp = client.post(
        "/webhooks/calendly",
        data=body_bytes,
        content_type="application/json",
        headers={"Calendly-Webhook-Signature": "t=1,v1=deadbeef"},
    )
    assert resp.status_code == 401
    assert db.get_booking_by_uuid("INVITEE_UUID") is None


def test_valid_signature_accepted(temp_db, fake_event_details, client, monkeypatch):
    monkeypatch.setattr(settings, "calendly_signing_secret", "test-secret")
    body = json.loads(FIXTURE_PATH.read_text())
    body_bytes = json.dumps(body).encode()
    sig_header = _sign(body_bytes, "test-secret")

    resp = client.post(
        "/webhooks/calendly",
        data=body_bytes,
        content_type="application/json",
        headers={"Calendly-Webhook-Signature": sig_header},
    )
    assert resp.status_code == 201


# ─── Leads sheet integration (Calendly path, requirement 5) ───────────

def test_invitee_created_appends_leads_sheet_row(temp_db, no_signature_check, fake_event_details, client, tmp_path, monkeypatch):
    from openpyxl import load_workbook
    from app import leads_sheet

    dest = tmp_path / "leads.xlsx"
    monkeypatch.setattr(settings, "leads_sheet_path", str(dest))
    monkeypatch.setattr(settings, "leads_sheet_enabled", True)

    body = json.loads(FIXTURE_PATH.read_text())
    resp = client.post("/webhooks/calendly", data=json.dumps(body), content_type="application/json")
    assert resp.status_code == 201

    booking = db.get_booking_by_uuid("INVITEE_UUID")
    assert booking["leads_sheet_synced_at"] is not None

    wb = load_workbook(dest)
    ws = wb.active
    data_row = [cell.value for cell in ws[2]]
    assert data_row[2]  # client name column populated
    assert "H-1B extension" in (data_row[3] or "")


def test_invitee_created_still_succeeds_when_leads_sheet_write_fails(
    temp_db, no_signature_check, fake_event_details, client, monkeypatch
):
    from app import leads_sheet

    def _boom(**kwargs):
        raise RuntimeError("disk full")

    monkeypatch.setattr(leads_sheet, "append_booking_row", _boom)

    body = json.loads(FIXTURE_PATH.read_text())
    resp = client.post("/webhooks/calendly", data=json.dumps(body), content_type="application/json")

    assert resp.status_code == 201
    booking = db.get_booking_by_uuid("INVITEE_UUID")
    assert booking is not None
    assert booking["leads_sheet_synced_at"] is None  # left for backfill to pick up


# ─── Fireflies webhook (requirements 2/3/4: capture + post-meeting notes) ─

FIREFLIES_SECRET = "test-fireflies-secret"


def _sign_fireflies(body_bytes: bytes, secret: str) -> str:
    return hmac.new(secret.encode(), body_bytes, hashlib.sha256).hexdigest()


def _make_synced_booking(join_url="https://teams.microsoft.com/l/meetup-join/abc"):
    client_id, _ = db.upsert_client("jane@example.com", "Jane Client")
    booking_id = db.create_booking(
        "evt-ff-1",
        client_id,
        "attorney@example.com",
        "2026-09-02T18:30:00+00:00",
        "Wants an H-1B extension",
        join_url,
        False,
    )
    return booking_id, client_id


def test_fireflies_webhook_rejects_bad_signature(temp_db, client, monkeypatch):
    monkeypatch.setattr(settings, "fireflies_webhook_secret", FIREFLIES_SECRET)
    body = {"meetingId": "m-1", "eventType": "Transcription completed"}
    body_bytes = json.dumps(body).encode()

    resp = client.post(
        "/webhooks/fireflies",
        data=body_bytes,
        content_type="application/json",
        headers={"x-hub-signature": "deadbeef"},
    )
    assert resp.status_code == 401


def test_fireflies_webhook_ignores_other_event_types(temp_db, client, monkeypatch):
    monkeypatch.setattr(settings, "fireflies_webhook_secret", "")
    body = {"meetingId": "m-1", "eventType": "Transcription started"}
    resp = client.post("/webhooks/fireflies", data=json.dumps(body), content_type="application/json")
    assert resp.status_code == 200
    assert resp.get_json()["status"] == "ignored"


def test_fireflies_webhook_accepts_meeting_transcribed_wording(temp_db, client, monkeypatch):
    """Fireflies' live webhook config screen offers "Meeting Transcribed"
    as an event name, not "Transcription completed" like their docs say
    — both must be accepted (case-insensitively) since it's unclear which
    one actually shows up in a real payload."""
    from app import fireflies_client

    monkeypatch.setattr(settings, "fireflies_webhook_secret", "")
    fake_transcript = {
        "id": "m-2",
        "meeting_link": None,
        "organizer_email": None,
        "sentences": [],
    }
    monkeypatch.setattr(fireflies_client, "get_transcript", lambda meeting_id: fake_transcript)

    body = {"meetingId": "m-2", "eventType": "meeting transcribed"}
    resp = client.post("/webhooks/fireflies", data=json.dumps(body), content_type="application/json")
    assert resp.status_code == 200
    assert resp.get_json()["status"] == "processed"


def test_fireflies_webhook_processes_matching_transcript(temp_db, client, monkeypatch):
    from app import fireflies_client, summarizer, teams_delivery

    monkeypatch.setattr(settings, "fireflies_webhook_secret", "")
    booking_id, client_id = _make_synced_booking()

    fake_transcript = {
        "id": "m-1",
        "meeting_link": "https://teams.microsoft.com/l/meetup-join/abc",
        "organizer_email": "attorney@example.com",
        "sentences": [{"speaker_name": "Jane Client", "text": "I need help with my H-1B."}],
    }
    monkeypatch.setattr(fireflies_client, "get_transcript", lambda meeting_id: fake_transcript)
    monkeypatch.setattr(summarizer, "summarize_transcript", lambda name, text: "- discussed H-1B extension")

    sent = []
    monkeypatch.setattr(
        teams_delivery,
        "send_manager_message",
        lambda manager_upn, title, lines: (sent.append((manager_upn, title, lines)), True)[1],
    )

    body = {"meetingId": "m-1", "eventType": "Transcription completed"}
    resp = client.post("/webhooks/fireflies", data=json.dumps(body), content_type="application/json")

    assert resp.status_code == 200
    assert resp.get_json()["status"] == "processed"

    booking = db.get_booking_by_uuid("evt-ff-1")
    assert booking["notes_synced_at"] is not None

    assert len(sent) == 1
    manager_upn, title, lines = sent[0]
    assert manager_upn == "attorney@example.com"
    assert "Jane Client" in title
    joined = "\n".join(lines)
    assert "H-1B extension" in joined
    assert "discussed H-1B extension" in joined


def test_fireflies_webhook_handles_unmatched_transcript_gracefully(temp_db, client, monkeypatch):
    from app import fireflies_client

    monkeypatch.setattr(settings, "fireflies_webhook_secret", "")
    fake_transcript = {
        "id": "m-unknown",
        "meeting_link": "https://teams.microsoft.com/l/meetup-join/nobody-booked-this",
        "organizer_email": "nobody@example.com",
        "sentences": [],
    }
    monkeypatch.setattr(fireflies_client, "get_transcript", lambda meeting_id: fake_transcript)

    body = {"meetingId": "m-unknown", "eventType": "Transcription completed"}
    resp = client.post("/webhooks/fireflies", data=json.dumps(body), content_type="application/json")

    # Not an error response — just nothing to correlate it to.
    assert resp.status_code == 200
    assert resp.get_json()["status"] == "processed"



# ─── Internal scheduled-job triggers (Cloud Run) ────────────────────

def test_internal_trigger_rejects_missing_secret(temp_db, client, monkeypatch):
    monkeypatch.setattr(settings, "internal_trigger_secret", "s3cret")
    resp = client.post("/internal/trigger/brief-scheduler")
    assert resp.status_code == 401


def test_internal_trigger_rejects_wrong_secret(temp_db, client, monkeypatch):
    monkeypatch.setattr(settings, "internal_trigger_secret", "s3cret")
    resp = client.post(
        "/internal/trigger/brief-scheduler",
        headers={"X-Internal-Trigger-Secret": "wrong"},
    )
    assert resp.status_code == 401


def test_internal_trigger_rejects_everything_when_secret_unset(temp_db, client, monkeypatch):
    monkeypatch.setattr(settings, "internal_trigger_secret", "")
    resp = client.post(
        "/internal/trigger/brief-scheduler",
        headers={"X-Internal-Trigger-Secret": ""},
    )
    assert resp.status_code == 401


def test_internal_trigger_brief_scheduler_runs_with_correct_secret(temp_db, client, monkeypatch):
    from app import brief_scheduler

    monkeypatch.setattr(settings, "internal_trigger_secret", "s3cret")
    calls = []
    monkeypatch.setattr(brief_scheduler, "run", lambda: calls.append(1))

    resp = client.post(
        "/internal/trigger/brief-scheduler",
        headers={"X-Internal-Trigger-Secret": "s3cret"},
    )

    assert resp.status_code == 200
    assert resp.get_json()["status"] == "ok"
    assert calls == [1]


def test_internal_trigger_live_catchup_runs_with_correct_secret(temp_db, client, monkeypatch):
    from app import live_catchup

    monkeypatch.setattr(settings, "internal_trigger_secret", "s3cret")
    calls = []
    monkeypatch.setattr(live_catchup, "run", lambda: calls.append(1))

    resp = client.post(
        "/internal/trigger/live-catchup",
        headers={"X-Internal-Trigger-Secret": "s3cret"},
    )

    assert resp.status_code == 200
    assert resp.get_json()["status"] == "ok"
    assert calls == [1]


def test_internal_trigger_returns_500_but_does_not_crash_on_failure(temp_db, client, monkeypatch):
    from app import brief_scheduler

    monkeypatch.setattr(settings, "internal_trigger_secret", "s3cret")

    def _boom():
        raise RuntimeError("groq is down")

    monkeypatch.setattr(brief_scheduler, "run", _boom)

    resp = client.post(
        "/internal/trigger/brief-scheduler",
        headers={"X-Internal-Trigger-Secret": "s3cret"},
    )

    assert resp.status_code == 500
